# anomaly_web/utils/file_handler.py
"""
File Upload/Download Handler
จัดการการอัพโหลด, ดาวน์โหลด, และเก็บ metadata ของไฟล์
"""

import os
import json
import uuid
import csv
import shutil
import tempfile
from datetime import date, datetime
from werkzeug.utils import secure_filename
from openpyxl import load_workbook
from filelock import FileLock


CSV_EXTENSIONS = {'.csv'}
EXCEL_EXTENSIONS = {'.xlsx', '.xlsm'}

class FileHandler:
    """จัดการไฟล์ input และ output พร้อม metadata"""
    
    def __init__(self, upload_folder, output_folder):
        self.upload_folder = upload_folder
        self.output_folder = output_folder
        self.metadata_file = os.path.join(upload_folder, '_metadata.json')
        self.metadata_lock_file = os.path.join(upload_folder, '_metadata.json.lock')
        self.output_metadata_file = os.path.join(output_folder, '_output_metadata.json')
        self.output_metadata_lock_file = os.path.join(output_folder, '_output_metadata.json.lock')
        
        # Create folders if not exist
        os.makedirs(upload_folder, exist_ok=True)
        os.makedirs(output_folder, exist_ok=True)
        
        # Load metadata
        self.metadata = self._load_metadata()
        self.output_metadata = self._load_output_metadata()
    
    def _load_metadata(self):
        """โหลด metadata ของไฟล์ที่ upload"""
        lock = FileLock(self.metadata_lock_file)
        with lock:
            if os.path.exists(self.metadata_file):
                with open(self.metadata_file, 'r', encoding='utf-8') as f:
                    return json.load(f)
        return {}
    
    def _save_metadata(self):
        """บันทึก metadata"""
        lock = FileLock(self.metadata_lock_file)
        with lock:
            with open(self.metadata_file, 'w', encoding='utf-8') as f:
                json.dump(self.metadata, f, ensure_ascii=False, indent=2)
    
    def _load_output_metadata(self):
        """โหลด metadata ของไฟล์ output"""
        lock = FileLock(self.output_metadata_lock_file)
        with lock:
            if os.path.exists(self.output_metadata_file):
                with open(self.output_metadata_file, 'r', encoding='utf-8') as f:
                    return json.load(f)
        return {}
    
    def _save_output_metadata(self):
        """บันทึก output metadata"""
        lock = FileLock(self.output_metadata_lock_file)
        with lock:
            with open(self.output_metadata_file, 'w', encoding='utf-8') as f:
                json.dump(self.output_metadata, f, ensure_ascii=False, indent=2)
    
    def save_upload(self, file, input_mode, description=''):
        """
        บันทึกไฟล์ที่ upload พร้อม metadata
        
        Returns:
            dict: ข้อมูลของไฟล์ที่บันทึก
        """
        # Generate unique ID
        file_id = str(uuid.uuid4())
        
        # Secure filename
        original_filename = secure_filename(file.filename)
        filename_without_ext = os.path.splitext(original_filename)[0]
        file_extension = os.path.splitext(original_filename)[1]
        
        # Save with unique name
        saved_filename = f"{file_id}_{original_filename}"
        filepath = os.path.join(self.upload_folder, saved_filename)
        file.save(filepath)
        
        # Get file info
        file_size = os.path.getsize(filepath)
        
        # Auto-generate tags
        tags = self._generate_tags(original_filename, input_mode)
        
        # Store metadata
        file_info = {
            'file_id': file_id,
            'original_filename': original_filename,
            'saved_filename': saved_filename,
            'filepath': filepath,
            'file_size': file_size,
            'file_size_mb': round(file_size / (1024 * 1024), 2),
            'input_mode': input_mode,
            'description': description,
            'tags': tags,
            'upload_time': datetime.now().isoformat(),
            'last_accessed': datetime.now().isoformat()
        }
        
        self.metadata[file_id] = file_info
        self._save_metadata()
        
        return file_info
    
    def save_uploads(self, files, input_mode, description='', progress_callback=None):
        """
        รวมหลายไฟล์ CSV/Excel ที่มี header เดียวกันเป็นไฟล์เดียว แล้วบันทึกเหมือน upload ปกติ

        Returns:
            dict: ข้อมูลของไฟล์ที่บันทึก (เหมือน save_upload)
        """
        files = [
            f for f in files
            if f.filename and not os.path.basename(f.filename).startswith('~$')
        ]
        if not files:
            raise ValueError("ไม่พบไฟล์ที่ใช้ได้สำหรับ upload")
        if len(files) == 1:
            self._notify_progress(
                progress_callback,
                status='saving',
                progress=50,
                message=f"Saving {files[0].filename}..."
            )
            file_info = self.save_upload(files[0], input_mode, description)
            self._notify_progress(
                progress_callback,
                status='completed',
                progress=100,
                message="Upload saved"
            )
            return file_info

        merge_types = {self._merge_type(f.filename) for f in files}
        if len(merge_types) != 1:
            raise ValueError("รวมหลายไฟล์ได้เฉพาะชนิดเดียวกันทั้งหมด: CSV หรือ Excel (.xlsx/.xlsm)")

        merge_type = merge_types.pop()
        if merge_type is None:
            bad_files = [f.filename for f in files]
            raise ValueError(f"ชนิดไฟล์ยังไม่รองรับสำหรับการรวม: {', '.join(bad_files)}")

        file_id = str(uuid.uuid4())
        base = os.path.splitext(secure_filename(files[0].filename))[0]
        original_filename = f"{base}_merged{len(files)}.csv"
        saved_filename = f"{file_id}_{original_filename}"
        filepath = os.path.join(self.upload_folder, saved_filename)

        self._notify_progress(
            progress_callback,
            status='merging',
            progress=5,
            message=f"Preparing to merge {len(files)} {merge_type.upper()} files..."
        )

        try:
            if merge_type == 'csv':
                self._merge_csv_uploads(files, filepath, progress_callback=progress_callback)
            else:
                self._merge_excel_uploads(files, filepath, progress_callback=progress_callback)
        except Exception:
            if os.path.exists(filepath):
                os.remove(filepath)
            raise

        self._notify_progress(
            progress_callback,
            status='saving_metadata',
            progress=95,
            message="Saving upload metadata..."
        )

        file_size = os.path.getsize(filepath)
        file_info = {
            'file_id': file_id,
            'original_filename': original_filename,
            'saved_filename': saved_filename,
            'filepath': filepath,
            'file_size': file_size,
            'file_size_mb': round(file_size / (1024 * 1024), 2),
            'input_mode': input_mode,
            'description': description,
            'tags': self._generate_tags(original_filename, input_mode) + ['merged'],
            'source_files': [f.filename for f in files],
            'upload_time': datetime.now().isoformat(),
            'last_accessed': datetime.now().isoformat()
        }

        self.metadata[file_id] = file_info
        self._save_metadata()

        self._notify_progress(
            progress_callback,
            status='completed',
            progress=100,
            message="Upload saved"
        )

        return file_info

    def _notify_progress(self, callback, **data):
        if callback:
            callback(data)

    def _merge_type(self, filename):
        ext = os.path.splitext(filename.lower())[1]
        if ext in CSV_EXTENSIONS:
            return 'csv'
        if ext in EXCEL_EXTENSIONS:
            return 'excel'
        return None

    def _merge_csv_uploads(self, files, filepath, progress_callback=None):
        # ponytail: ต่อไฟล์ระดับ byte (ตัด header ของไฟล์ที่ 2 เป็นต้นไป)
        # ไม่ใช้ pandas -> ไม่กิน RAM, ไม่ต้องเดา encoding/delimiter
        header = None
        total_files = len(files)
        with open(filepath, 'wb') as out:
            for idx, f in enumerate(files, 1):
                self._notify_progress(
                    progress_callback,
                    status='merging',
                    progress=round(5 + ((idx - 1) / total_files * 90)),
                    message=f"Merging CSV file {idx}/{total_files}: {f.filename}"
                )
                f.stream.seek(0)
                first_line = f.stream.readline()
                if header is None:
                    header = first_line
                    out.write(first_line)
                elif first_line.strip() != header.strip():
                    raise ValueError(
                        f"หัวคอลัมน์ของ {f.filename} ไม่ตรงกับไฟล์แรก ({files[0].filename})"
                    )
                last = b'\n'
                for chunk in iter(lambda: f.stream.read(1 << 20), b''):
                    out.write(chunk)
                    last = chunk[-1:]
                if last != b'\n':
                    out.write(b'\n')
                self._notify_progress(
                    progress_callback,
                    status='merging',
                    progress=round(5 + (idx / total_files * 90)),
                    message=f"Merged CSV file {idx}/{total_files}: {f.filename}"
                )

    def _merge_excel_uploads(self, files, filepath, progress_callback=None):
        header = None
        total_files = len(files)
        with open(filepath, 'w', encoding='utf-8-sig', newline='') as out:
            writer = csv.writer(out)
            for idx, f in enumerate(files, 1):
                file_start_progress = 5 + ((idx - 1) / total_files * 90)
                file_end_progress = 5 + (idx / total_files * 90)
                self._notify_progress(
                    progress_callback,
                    status='merging',
                    progress=round(file_start_progress),
                    message=f"Preparing Excel file {idx}/{total_files}: {f.filename}"
                )
                f.stream.seek(0)
                copy_done_progress = file_start_progress + ((file_end_progress - file_start_progress) * 0.25)
                temp_path = self._copy_upload_to_temp_file(
                    f,
                    progress_callback=progress_callback,
                    progress_start=file_start_progress,
                    progress_end=copy_done_progress,
                    file_index=idx,
                    total_files=total_files
                )
                wb = None
                try:
                    try:
                        wb = load_workbook(temp_path, read_only=True, data_only=True)
                    except Exception as e:
                        raise ValueError(f"อ่านไฟล์ Excel {f.filename} ไม่ได้: {e}") from e

                    ws = wb.active
                    rows = ws.iter_rows(values_only=True)
                    try:
                        current_header = next(rows)
                    except StopIteration:
                        raise ValueError(f"ไฟล์ {f.filename} ไม่มีข้อมูล")

                    current_header = [self._excel_cell_to_csv(value).strip() for value in current_header]
                    if header is None:
                        header = current_header
                        writer.writerow(header)
                    elif current_header != header:
                        raise ValueError(
                            f"หัวคอลัมน์ของ {f.filename} ไม่ตรงกับไฟล์แรก ({files[0].filename})"
                        )

                    max_rows = max((ws.max_row or 1) - 1, 1)
                    for row_idx, row in enumerate(rows, 1):
                        if row is None or all(value is None for value in row):
                            continue
                        writer.writerow([self._excel_cell_to_csv(value) for value in row])
                        if row_idx % 10000 == 0:
                            row_ratio = min(row_idx / max_rows, 1)
                            self._notify_progress(
                                progress_callback,
                                status='merging',
                                progress=round(copy_done_progress + ((file_end_progress - copy_done_progress) * row_ratio)),
                                message=(
                                    f"Merging Excel file {idx}/{total_files}: {f.filename} "
                                    f"({row_idx:,}/{max_rows:,} rows)"
                                )
                            )
                finally:
                    if wb is not None:
                        wb.close()
                    if os.path.exists(temp_path):
                        os.remove(temp_path)
                self._notify_progress(
                    progress_callback,
                    status='merging',
                    progress=round(file_end_progress),
                    message=f"Merged Excel file {idx}/{total_files}: {f.filename}"
                )

    def _copy_upload_to_temp_file(
        self,
        file,
        progress_callback=None,
        progress_start=0,
        progress_end=0,
        file_index=1,
        total_files=1
    ):
        suffix = os.path.splitext(file.filename)[1]
        total_size = self._stream_size(file.stream)
        temp = tempfile.NamedTemporaryFile(
            mode='wb',
            suffix=suffix,
            prefix='upload_merge_',
            dir=self.upload_folder,
            delete=False
        )
        try:
            with temp:
                copied = 0
                last_reported = 0
                for chunk in iter(lambda: file.stream.read(1 << 20), b''):
                    temp.write(chunk)
                    copied += len(chunk)
                    if copied - last_reported >= 20 * 1024 * 1024:
                        last_reported = copied
                        if total_size:
                            ratio = min(copied / total_size, 1)
                            progress = progress_start + ((progress_end - progress_start) * ratio)
                            detail = f"{copied / 1024 / 1024:,.0f}/{total_size / 1024 / 1024:,.0f} MB"
                        else:
                            progress = progress_start
                            detail = f"{copied / 1024 / 1024:,.0f} MB"
                        self._notify_progress(
                            progress_callback,
                            status='merging',
                            progress=round(progress),
                            message=f"Copying Excel file {file_index}/{total_files}: {file.filename} ({detail})"
                        )
            return temp.name
        except Exception:
            if os.path.exists(temp.name):
                os.remove(temp.name)
            raise

    def _stream_size(self, stream):
        try:
            pos = stream.tell()
            stream.seek(0, os.SEEK_END)
            size = stream.tell()
            stream.seek(pos)
            return size
        except Exception:
            return None

    def _excel_cell_to_csv(self, value):
        if value is None:
            return ''
        if isinstance(value, datetime):
            return value.isoformat(sep=' ')
        if isinstance(value, date):
            return value.isoformat()
        return str(value)

    def update_upload_metadata(self, file_id, description=None, tags=None):
        """
        อัพเดท metadata ของไฟล์ที่ upload
        """
        if file_id not in self.metadata:
            return False
            
        if description is not None:
            self.metadata[file_id]['description'] = description
        
        if tags is not None:
            # Ensure tags are a list of unique strings
            self.metadata[file_id]['tags'] = sorted(list(set(tags)))
            
        self.metadata[file_id]['last_accessed'] = datetime.now().isoformat()
        self._save_metadata()
        return True
    
    def _generate_tags(self, filename, input_mode):
        """สร้าง tags อัตโนมัติจากชื่อไฟล์และ mode"""
        tags = [input_mode]
        
        # Extract year from filename
        import re
        year_match = re.search(r'20\d{2}', filename)
        if year_match:
            tags.append(year_match.group())
        
        # Extract common keywords
        keywords = ['expense', 'revenue', 'report', 'audit', 'nt', 'cost']
        filename_lower = filename.lower()
        for keyword in keywords:
            if keyword in filename_lower:
                tags.append(keyword)
        
        return list(set(tags))  # Remove duplicates
    
    def get_file_info(self, file_id):
        """ดึงข้อมูลไฟล์จาก ID"""
        file_info = self.metadata.get(file_id)
        if file_info:
            # Update last accessed
            file_info['last_accessed'] = datetime.now().isoformat()
            self.metadata[file_id] = file_info
            self._save_metadata()
        return file_info
    
    def list_uploads(self, limit=50, sort_by='upload_time', reverse=True):
        """
        แสดงรายการไฟล์ที่ upload ทั้งหมด
        
        Args:
            limit: จำนวนไฟล์สูงสุด
            sort_by: เรียงตาม field ใด
            reverse: เรียงจากมากไปน้อย
        """
        files = list(self.metadata.values())
        files.sort(key=lambda x: x.get(sort_by, ''), reverse=reverse)
        return files[:limit]
    
    def delete_upload(self, file_id):
        """ลบไฟล์ที่ upload"""
        file_info = self.metadata.get(file_id)
        if file_info:
            # Delete file
            if os.path.exists(file_info['filepath']):
                os.remove(file_info['filepath'])
            
            # Delete metadata
            del self.metadata[file_id]
            self._save_metadata()
            
            return True
        return False
    
    def save_output_info(self, file_id, output_filename, output_path, config, result):
        """บันทึกข้อมูล output file"""
        output_id = str(uuid.uuid4())

        output_info = {
            'output_id': output_id,
            'input_file_id': file_id,
            'filename': output_filename,
            'filepath': output_path,
            'file_size': os.path.getsize(output_path),
            'file_size_mb': round(os.path.getsize(output_path) / (1024 * 1024), 2),
            'created_time': datetime.now().isoformat(),
            'config_summary': self._summarize_config(config),
            'result_summary': result,
            'tags': self._generate_output_tags(config, result)
        }

        self.output_metadata[output_id] = output_info
        self._save_output_metadata()

        return output_info
    
    def _summarize_config(self, config):
        """สร้างสรุป config แบบสั้น"""
        return {
            'input_mode': config.get('input_mode'),
            'time_series_enabled': config.get('run_time_series_analysis', False),
            'peer_group_enabled': config.get('run_peer_group_analysis', False),
            'dimensions': ', '.join(config.get('crosstab_dimensions', []))
        }

    def _generate_output_tags(self, config, result):
        """สร้าง tags สำหรับ output file จาก config และผลลัพธ์"""
        tags = []

        # Input mode
        input_mode = config.get('input_mode', 'long')
        tags.append(input_mode)

        # Analysis types
        if config.get('run_time_series_analysis'):
            tags.append('time_series')
            # Add window size if specified
            window = config.get('audit_ts_window')
            if window:
                tags.append(f'window_{window}')

        if config.get('run_peer_group_analysis'):
            tags.append('peer_group')

        if config.get('run_crosstab_report'):
            tags.append('crosstab')

        # Dimensions used
        dimensions = config.get('crosstab_dimensions', [])
        if dimensions:
            # Add count of dimensions
            tags.append(f'{len(dimensions)}_dims')
            # Add abbreviated dimension names (first 3 only to avoid too many tags)
            for dim in dimensions[:3]:
                # Abbreviate dimension name
                dim_abbr = dim.replace('_', '').lower()[:10]
                tags.append(dim_abbr)

        # Target column
        target_col = config.get('target_col')
        if target_col:
            # Abbreviate target column name
            target_abbr = target_col.replace('_', '').lower()[:10]
            tags.append(target_abbr)

        # Result-based tags
        if result:
            ts_count = result.get('ts_anomalies', 0)
            peer_count = result.get('peer_anomalies', 0)

            # Add anomaly count ranges
            if ts_count > 0:
                if ts_count > 1000:
                    tags.append('high_anomalies')
                elif ts_count > 100:
                    tags.append('medium_anomalies')
                else:
                    tags.append('low_anomalies')

        # Remove duplicates and limit to 10 tags
        tags = list(dict.fromkeys(tags))[:10]

        return tags
    
    def get_output_info(self, output_id):
        """ดึงข้อมูล output file"""
        return self.output_metadata.get(output_id)
    
    def list_outputs(self, limit=50, sort_by='created_time', reverse=True):
        """แสดงรายการ output ทั้งหมด"""
        outputs = list(self.output_metadata.values())
        outputs.sort(key=lambda x: x.get(sort_by, ''), reverse=reverse)
        return outputs[:limit]
    
    def delete_output(self, output_id):
        """ลบ output file"""
        output_info = self.output_metadata.get(output_id)
        if output_info:
            # Delete file
            if os.path.exists(output_info['filepath']):
                os.remove(output_info['filepath'])
            
            # Delete metadata
            del self.output_metadata[output_id]
            self._save_output_metadata()
            
            return True
        return False
