"""เช็ค data preparation/report: python3 test_data_preparation.py"""
import os
import shutil
import tempfile

import pandas as pd
from openpyxl import load_workbook

import utils.anomaly_reporter as reporter_mod
from utils.audit_runner import AuditRunner
from utils.anomaly_engine import CrosstabGenerator
from utils.anomaly_reporter import ExcelReporter
from utils.data_analyzer import DataAnalyzer


def main():
    tmp = tempfile.mkdtemp()

    runner = AuditRunner(tmp)
    df = pd.DataFrame({
        "DATE": ["01.02.2026", "02.02.2026", "bad-date", "03.03.2026", ""],
        "VALUE": ["1,000", "oops", "100", "(50)", "999999"],
        "GROUP": ["A", "A", "A", "A", ""],
    })
    clean = runner._prepare_data(df, {
        "date_column": "DATE",
        "date_parse_mode": "auto",
        "date_grain": "month",
        "bad_date_policy": "drop",
        "target_col": "VALUE",
        "bad_value_policy": "drop",
        "value_multiplier": -1,
        "value_add": 5,
        "crosstab_dimensions": ["GROUP"],
    })
    assert clean["__date_col__"].dt.strftime("%Y-%m-%d").tolist() == ["2026-02-01", "2026-03-01"]
    assert clean["VALUE"].tolist() == [-995, 55]

    analysis = DataAnalyzer().analyze_dataframe(pd.DataFrame({
        "ว/ทเอกสาร": ["01.01.2016", "01.02.2017", "01.03.2018", "01.04.2019"],
        "Pstng Date": ["01.08.2026", "02.08.2026", "03.08.2026", "04.08.2026"],
        "บัญชี": [1000, 2000, 3000, 4000],
        "เลขเอกสาร": [111, 222, 333, 444],
        "จำนวนเงินในสกุลในปท.": ["1,000", "2,000", "oops", "3,000"],
        "SERVICE_GROUP_SEQ": [1, 1, 2, 2],
    }), input_mode="long")
    assert analysis["recommendations"]["date_column"] == "Pstng Date"
    value_names = [col["name"] for col in analysis["recommendations"]["value_columns"]]
    assert "จำนวนเงินในสกุลในปท." in value_names, value_names
    assert "บัญชี" not in value_names, value_names
    assert "เลขเอกสาร" not in value_names, value_names

    report = CrosstabGenerator(pd.DataFrame({
        "GROUP": ["A", "A", "A", "A"],
        "__date_col__": pd.to_datetime(["2026-01-01", "2026-01-20", "2026-02-01", "2026-03-01"]),
        "VALUE": [60, 40, 80, 100],
    }), min_history=1, anomaly_settings={"date_grain": "month"}).create_report("VALUE", "__date_col__", ["GROUP"])
    assert [col for col in report.columns if col.startswith("2026-")] == ["2026-01", "2026-02", "2026-03"]
    row = report.iloc[0]
    assert row["PREVIOUS_VALUE"] == 80
    assert row["DIFF_PREVIOUS"] == 20
    assert row["PCT_DIFF_PREVIOUS"] == 25

    output = os.path.join(tmp, "split.xlsx")
    old_limit = reporter_mod.EXCEL_MAX_ROWS
    reporter_mod.EXCEL_MAX_ROWS = 4
    try:
        reporter = ExcelReporter(output)
        reporter.add_audit_log_sheet(pd.DataFrame({"A": range(7)}), "Long_Log", ["A"])
        reporter.save()
        wb = load_workbook(output, read_only=True)
        try:
            assert wb.sheetnames == ["Long_Log_1", "Long_Log_2", "Long_Log_3"], wb.sheetnames
        finally:
            wb.close()
    finally:
        reporter_mod.EXCEL_MAX_ROWS = old_limit
        shutil.rmtree(tmp)

    print("OK")


if __name__ == "__main__":
    main()
